#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
템플릿(Template.pptx / Template2.pptx)에 학생 이름+내용 데이터를 자동으로 채워 넣는 스크립트.
- 슬라이드 1장당 (표 개수 x 3칸)까지 채움. Template.pptx는 표 4개라 12명,
  Template2.pptx는 표 3개라 9명.
- 수용 인원을 초과하면 슬라이드를 자동으로 복제하여 다음 페이지에 이어서 채움
- 모자라면 남은 칸은 빈칸으로 둠 ({{이름}}/{{내용}} 마커만 제거)
"""
import os
import re
import sys
import shutil
import subprocess
from pathlib import Path
from pptx import Presentation

# 표 하나 안에서 데이터를 채울 시작 열번호 (두 템플릿 모두 동일)
COL_STARTS = (0, 7, 14)

NAME_ROW = 0
CONTENT_ROW = 3


def slot_order(slide):
    """슬라이드에서 (표, 시작 열번호) 슬롯 목록을 채울 순서대로 반환.

    표 이름(예: '표 1')은 템플릿마다 다르고 PowerPoint에서 재저장하면 바뀌기도 하므로
    이름 대신 화면상 위치(위 -> 아래, 왼쪽 -> 오른쪽)로 순서를 정한다.
    """
    tables = sorted(
        (shape for shape in slide.shapes if shape.has_table),
        key=lambda shape: (shape.top, shape.left),
    )
    return [(shape.table, col) for shape in tables for col in COL_STARTS]


def slots_per_slide(template_path):
    """템플릿 첫 슬라이드가 수용하는 인원 수."""
    return len(slot_order(Presentation(template_path).slides[0]))


def parse_records(raw_text):
    """'■ 이름 내용...' 형태의 레코드가 여러 개 붙은 텍스트를 리스트로 분리.

    API는 Make.com에서 이미 구조화된 JSON 배열을 받으므로 보통 이 함수는 쓰이지 않는다.
    텍스트 블록을 그대로 보내는 경우를 대비한 폴백 파서로만 남겨둠.
    """
    chunks = re.split(r"(?=■)", raw_text.strip())
    records = []
    for chunk in chunks:
        chunk = chunk.strip()
        if not chunk:
            continue
        chunk = chunk.lstrip("■").strip()
        m = re.match(r"^(\S+)\s+(.*)$", chunk, re.DOTALL)
        if not m:
            continue
        name, content = m.group(1), m.group(2).strip()
        records.append({"name": name, "content": content})
    return records


def set_cell_text(cell, new_text):
    """플레이스홀더가 여러 run으로 쪼개져 있어도 첫 run의 서식을 유지하며 텍스트 교체."""
    tf = cell.text_frame
    para = tf.paragraphs[0]
    runs = para.runs
    if not runs:
        para.text = new_text
        return
    runs[0].text = new_text
    for r in runs[1:]:
        r.text = ""
    # 혹시 플레이스홀더 아래 추가 문단이 있으면 텍스트 비우기 (서식은 유지)
    for extra_p in tf.paragraphs[1:]:
        for r in extra_p.runs:
            r.text = ""


def fill_slide(slide, records_chunk):
    """records_chunk: 슬라이드 수용 인원 이하의 dict({'name':..,'content':..}) 목록"""
    slots = slot_order(slide)
    padded = list(records_chunk) + [None] * (len(slots) - len(records_chunk))

    for (table, col_start), record in zip(slots, padded):
        name_cell = table.cell(NAME_ROW, col_start)
        content_cell = table.cell(CONTENT_ROW, col_start)
        if record is None:
            set_cell_text(name_cell, "")
            set_cell_text(content_cell, "")
        else:
            set_cell_text(name_cell, record["name"])
            set_cell_text(content_cell, record["content"])


def build_pptx(template_path, records, output_path, scripts_dir=None, work_dir=None):
    """records가 비어 있으면 빈칸투성이 1장을 생성한다.

    work_dir: 슬라이드 복제 작업용 임시 디렉터리. 동시 요청이 서로의 임시 파일을
    덮어쓰지 않도록 호출자(API 서버)가 요청마다 고유한 디렉터리를 넘겨야 한다.
    """
    scripts_dir = Path(scripts_dir) if scripts_dir else Path(__file__).parent / "pptx_scripts"
    work_dir = Path(work_dir) if work_dir else Path(output_path).parent
    work_dir.mkdir(parents=True, exist_ok=True)

    per_slide = slots_per_slide(template_path)
    n_slides_needed = max(1, -(-len(records) // per_slide))  # ceil division

    # 1) 필요한 만큼 슬라이드 복제 (템플릿은 slide1.xml 한 장짜리)
    work_path = work_dir / "_work.pptx"
    shutil.copy(template_path, work_path)
    env = {**os.environ, "PYTHONIOENCODING": "utf-8"}
    for _ in range(1, n_slides_needed):
        result = subprocess.run(
            [sys.executable, str(scripts_dir / "add_slide.py"), str(work_path),
             "slide1.xml", "-o", str(work_path)],
            capture_output=True, text=True, encoding="utf-8", env=env
        )
        if result.returncode != 0:
            raise RuntimeError(f"add_slide.py failed: {result.stderr}")

    # 2) 각 슬라이드에 데이터 채우기
    prs = Presentation(work_path)
    chunks = [records[i:i + per_slide] for i in range(0, len(records), per_slide)]
    if not chunks:
        chunks = [[]]
    for slide, chunk in zip(prs.slides, chunks):
        fill_slide(slide, chunk)

    prs.save(output_path)
    return n_slides_needed


if __name__ == "__main__":
    sample_text = """■ 배장우 수업 시간 동안 양호한 자세와 태도를 유지하며 수업에 임하는 모습을 보였습니다. 타인에게 방해가 되는 행동은 관찰되지 않았으며, 이는 주변을 배려하며 안정적으로 수업에 참여하는 태도를 잘 보여주고 있습니다. 전반적으로 차분하고 무난한 수업 참여 자세를 유지한 것으로 확인되었습니다"""

    records = parse_records(sample_text)
    print("파싱된 레코드:", records)

    n = build_pptx("Template.pptx", records, "test_1record.pptx")
    print(f"생성된 슬라이드 수: {n}")
